"""Gera docs/artigo_consciencia_loop_autopreditivo.docx a partir de artigo.md.

Cobre o subconjunto de Markdown usado no artigo: títulos, parágrafos, listas, citação em bloco, bloco de código,
negrito, itálico, código em linha e links. Comentários HTML são ignorados.

uso: python3 scripts/gerar_docx.py
requer: python-docx
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "artigo.md"
OUT = ROOT / "docs" / "artigo_consciencia_loop_autopreditivo.docx"

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
        elif part.startswith("["):
            par.add_run(re.match(r"\[([^\]]+)\]", part).group(1))
        elif part.startswith("*"):
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def main():
    text = re.sub(r"<!--.*?-->", "", SRC.read_text(encoding="utf-8"), flags=re.S)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    lines = text.splitlines()
    i, buf = 0, []

    def flush():
        if buf:
            # a line ending in two spaces is a Markdown hard break
            joined = "".join(ln.strip() + ("\n" if ln.endswith("  ") else " ") for ln in buf)
            add_inline(doc.add_paragraph(), joined.strip())
            buf.clear()

    first_title = True
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s.startswith("```"):
            flush()
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(code))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif s.startswith("#"):
            flush()
            level = len(s) - len(s.lstrip("#"))
            h = doc.add_heading(level=0 if first_title else min(level, 3))
            add_inline(h, s.lstrip("#").strip())
            if first_title:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first_title = False
        elif s == "---" or not s:
            flush()
        elif s.startswith(">"):
            flush()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            add_inline(doc.add_paragraph(style="Intense Quote"), " ".join(quote))
            continue
        elif re.match(r"^(- |\d+\. )", s):
            flush()
            # numbered items keep their literal number: python-docx "List Number" would continue one counter
            # across every list of the document
            numbered = bool(re.match(r"^\d+\. ", s))
            item = [s if numbered else s[2:]]
            i += 1
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip():
                item.append(lines[i].strip())
                i += 1
            p = doc.add_paragraph() if numbered else doc.add_paragraph(style="List Bullet")
            add_inline(p, " ".join(item))
            continue
        else:
            buf.append(ln)
        i += 1
    flush()
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(f"gerado: {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
